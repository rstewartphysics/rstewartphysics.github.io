"""Cut a drawing out of a shipped deck PDF as inline SVG, themed for the web page.

    svg = cut("L1 - Microcontrollers.pdf", "The Transistor Driver", (4.78, 1.4, 12.71, 7.1),
              "Wiring: ...")

The figure is the deck's own vector drawing (same symbols, words and pins as the slide), so the
page and the lesson cannot drift apart. Recoloured for both appearances: ink -> currentColor,
the deck's accent text -> var(--eng-ink-accent), card and slide grounds -> none, white part
bodies and picture cards -> none. Fonts inherit from the page. Images (cue icons, photos) are dropped,
and the viewBox is trimmed to what is drawn inside the cut.

Each cut is saved in figures/, which is committed, so the page builds from the repo without the
decks (CI has no OneDrive). build_page.py --recut clears figures/ and cuts again from the decks.
"""
import glob
import hashlib
import os
import re
from html import escape

DECKS = os.path.expanduser("~/claude-work/n5-engineering/b4b/decks")
CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")
_recut = False
COLOURS = {"#2b2b2b": "currentColor", "#965719": "var(--eng-ink-accent)",
           "#ece5d8": "none", "#f7f4ee": "none", "#ffffff": "none"}
_titles = {}


def recut():
    global _recut
    _recut = True
    os.makedirs(CACHE, exist_ok=True)
    for f in glob.glob(os.path.join(CACHE, "*.svg")):
        os.remove(f)


def _cached(pdf, title, box, alt, max_w):
    key = hashlib.sha1(repr((pdf, title, box, alt, max_w)).encode()).hexdigest()[:8]
    slug = re.sub(r"[^a-z0-9]+", "-", f"{pdf[:2]} {title}".lower()).strip("-")
    return os.path.join(CACHE, f"{slug}-{key}.svg")


def _page(pdf, title):
    from pptx import Presentation
    if pdf not in _titles:
        t = {}
        for i, s in enumerate(Presentation(os.path.join(DECKS, pdf[:-4] + ".pptx")).slides):
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    t.setdefault(sh.text_frame.text.strip(), i)
                    break
        _titles[pdf] = t
    return _titles[pdf][title]


def _trim(page, crop, pad=6):
    import fitz
    """Shrink the cut to what is drawn in it: every line, shape and word inside the box, but not
    the card behind them (a filled rectangle nearly the size of the box)."""
    area = crop.width * crop.height
    found = fitz.Rect()
    for d in page.get_drawings():
        r = d["rect"]
        if not crop.contains(r) or r.width * r.height > 0.6 * area:
            continue
        found |= r
    for b in page.get_text("dict")["blocks"]:
        r = fitz.Rect(b["bbox"])
        if crop.contains(r):
            found |= r
    if found.is_empty:
        return crop
    return (found + (-pad, -pad, pad, pad)) & crop


def cut(pdf, title, box, alt, max_w=None):
    path = _cached(pdf, title, box, alt, max_w)
    if not _recut:
        if not os.path.exists(path):
            raise SystemExit(f"no saved cut for {pdf} / {title}: run build_page.py --recut on the Mac")
        with open(path, encoding="utf-8") as f:
            return f.read()
    svg = _cut(pdf, title, box, alt, max_w)
    with open(path, "w", encoding="utf-8") as f:
        f.write(svg)
    return svg


def _cut(pdf, title, box, alt, max_w):
    import fitz
    doc = fitz.open(os.path.join(DECKS, pdf))
    page = doc[_page(pdf, title)]
    x0, y0, x1, y1 = box
    page.set_cropbox(_trim(page, fitz.Rect(x0 * 72, y0 * 72, x1 * 72, y1 * 72)))
    svg = page.get_svg_image(text_as_path=0)
    svg = re.sub(r"<image\b[^>]*/>", "", svg, flags=re.S)
    svg = re.sub(r"<image\b.*?</image>", "", svg, flags=re.S)
    for old, new in COLOURS.items():
        svg = re.sub(old, new, svg, flags=re.I)
    unknown = set(re.findall(r'(?:fill|stroke)="(#[0-9a-fA-F]{6})"', svg))
    if unknown:
        raise ValueError(f"{pdf} / {title}: colours with no theme mapping {unknown}")
    svg = re.sub(r'\sfont-family="[^"]*"', "", svg)
    svg = re.sub(r'<\?xml[^>]*>\s*', "", svg)
    w = re.search(r'width="([\d.]+)"', svg).group(1)
    h = re.search(r'height="([\d.]+)"', svg).group(1)
    style = f' style="max-width:{max_w}px"' if max_w else ""
    svg = re.sub(r"<svg\b[^>]*>",
                 f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
                 f'class="deckfig" viewBox="0 0 {w} {h}" role="img" aria-label="{escape(alt)}" '
                 f'focusable="false"{style}><title>{escape(alt)}</title>', svg, count=1)
    # ids are per-file (clip_1, mask_3 …): make them unique on a page holding many figures
    tag = re.sub(r"\W", "", f"{pdf[:2]}{_page(pdf, title)}_{int(x0 * 10)}")
    svg = re.sub(r'\bid="([^"]+)"', rf'id="{tag}_\1"', svg)
    svg = re.sub(r'url\(#([^)]+)\)', rf'url(#{tag}_\1)', svg)
    svg = re.sub(r'xlink:href="#([^"]+)"', rf'xlink:href="#{tag}_\1"', svg)
    return svg
